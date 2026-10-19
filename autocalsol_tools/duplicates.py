from __future__ import annotations

import csv
from collections import defaultdict
from dataclasses import dataclass, field
from itertools import chain
from pathlib import Path
from typing import Iterable, Iterator

import pandas as pd
from openpyxl import load_workbook

from .normalization import normalized_text


SUPPORTED_TABLES = {".csv", ".xlsx", ".xlsm", ".xls", ".ods"}


@dataclass(frozen=True, slots=True)
class DuplicateOptions:
    selected_columns: tuple[str, ...] | None = None
    trim_values: bool = True
    lowercase_values: bool = True
    ignore_zero_values: bool = True
    ignore_unnamed_columns: bool = True


@dataclass(frozen=True, slots=True)
class DuplicateOccurrence:
    value: str
    column: str
    source: str
    sheet: str
    row_number: int
    duplicate_type: str
    row_content: str = ""


@dataclass(slots=True)
class DuplicateSummary:
    file_count: int = 0
    table_count: int = 0
    group_count: int = 0
    internal_count: int = 0
    external_count: int = 0
    occurrence_count: int = 0
    scanned_value_count: int = 0
    ignored_value_count: int = 0


@dataclass(slots=True)
class DuplicateAnalysis:
    occurrences: list[DuplicateOccurrence]
    summary: DuplicateSummary
    errors: list[str] = field(default_factory=list)


def supported_files_in_folder(folder: str | Path, recursive: bool = True) -> list[Path]:
    root = Path(folder)
    iterator = root.rglob("*") if recursive else root.glob("*")
    return sorted(path for path in iterator if path.is_file() and path.suffix.casefold() in SUPPORTED_TABLES)


def get_available_columns(paths: Iterable[str | Path]) -> tuple[list[str], list[str]]:
    columns: list[str] = []
    known: set[str] = set()
    errors: list[str] = []
    for raw_path in paths:
        path = Path(raw_path)
        try:
            for _, headers, _ in _iter_tables(path, headers_only=True):
                for header in headers:
                    name = str(header).strip()
                    key = normalized_text(name)
                    if name and key and key not in known:
                        known.add(key)
                        columns.append(name)
        except Exception as exc:
            errors.append(f"{path.name} : {exc}")
    return columns, errors


def analyze_duplicates(
    paths: Iterable[str | Path],
    options: DuplicateOptions | None = None,
) -> DuplicateAnalysis:
    settings = options or DuplicateOptions()
    selected = (
        {normalized_text(column) for column in settings.selected_columns}
        if settings.selected_columns
        else None
    )
    grouped: dict[tuple[str, str], list[tuple[str, str, int, str, str]]] = defaultdict(list)
    summary = DuplicateSummary()
    errors: list[str] = []
    unique_paths = [Path(path) for path in dict.fromkeys(str(Path(path)) for path in paths)]
    summary.file_count = len(unique_paths)

    for path in unique_paths:
        try:
            for sheet, headers, rows in _iter_tables(path):
                summary.table_count += 1
                usable = [
                    (index, str(header).strip(), normalized_text(header))
                    for index, header in enumerate(headers)
                    if str(header).strip()
                    and (not settings.ignore_unnamed_columns or not normalized_text(header).startswith("unnamed"))
                    and (selected is None or normalized_text(header) in selected)
                ]
                for row_number, row in rows:
                    row_content = " | ".join(
                        f"{headers[index]}={row[index]}"
                        for index in range(min(len(headers), len(row)))
                        if row[index] is not None and str(row[index]).strip()
                    )
                    for index, header, header_key in usable:
                        summary.scanned_value_count += 1
                        value = row[index] if index < len(row) else None
                        display, normalized = _normalize_value(value, settings)
                        if normalized is None:
                            summary.ignored_value_count += 1
                            continue
                        grouped[(header_key, normalized)].append(
                            (str(path), sheet, row_number, display, row_content)
                        )
        except Exception as exc:
            errors.append(f"{path.name} : {exc}")

    occurrences: list[DuplicateOccurrence] = []
    for (column_key, _), records in grouped.items():
        if len(records) < 2:
            continue
        sources = {record[0] for record in records}
        duplicate_type = "EXTERNE" if len(sources) > 1 else "INTERNE"
        header = _display_column(column_key, settings.selected_columns)
        if duplicate_type == "EXTERNE":
            summary.external_count += 1
        else:
            summary.internal_count += 1
        summary.group_count += 1
        for source, sheet, row_number, display, row_content in records:
            occurrences.append(
                DuplicateOccurrence(display, header, source, sheet, row_number, duplicate_type, row_content)
            )
    occurrences.sort(key=lambda item: (item.column.casefold(), item.value.casefold(), item.source, item.sheet, item.row_number))
    summary.occurrence_count = len(occurrences)
    return DuplicateAnalysis(occurrences, summary, errors)


def find_duplicates(
    paths: Iterable[str | Path], *, selected_column: str | None = None
) -> list[DuplicateOccurrence]:
    selected = (selected_column,) if selected_column else None
    return analyze_duplicates(paths, DuplicateOptions(selected_columns=selected)).occurrences


def export_duplicates(
    results: DuplicateAnalysis | list[DuplicateOccurrence], output_path: str | Path
) -> Path:
    occurrences = results.occurrences if isinstance(results, DuplicateAnalysis) else results
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = [
        {
            "Colonne": item.column,
            "Valeur": item.value,
            "Type": item.duplicate_type,
            "Fichier": item.source,
            "Feuille": item.sheet,
            "Ligne": item.row_number,
            "Ligne complète": item.row_content,
        }
        for item in occurrences
    ]
    frame = pd.DataFrame(rows, columns=["Colonne", "Valeur", "Type", "Fichier", "Feuille", "Ligne", "Ligne complète"])
    if path.suffix.casefold() == ".xlsx":
        frame.to_excel(path, index=False)
    else:
        frame.to_csv(path, index=False, sep=";", encoding="utf-8-sig")
    return path


def _normalize_value(value: object, options: DuplicateOptions) -> tuple[str, str | None]:
    if value is None or value is pd.NA or value is pd.NaT:
        return "", None
    if isinstance(value, float) and value != value:
        return "", None
    display = str(value)
    if options.trim_values:
        display = display.strip()
    normalized = display.casefold() if options.lowercase_values else display
    if not normalized:
        return display, None
    if options.ignore_zero_values and normalized.replace(",", ".") in {"0", "0.0", "0.00"}:
        return display, None
    return display, normalized


def _display_column(key: str, selected: tuple[str, ...] | None) -> str:
    if selected:
        for column in selected:
            if normalized_text(column) == key:
                return column
    return key


def _iter_tables(path: Path, headers_only: bool = False) -> Iterator[tuple[str, list[str], Iterator[tuple[int, tuple[object, ...]]]]]:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        encoding = _detect_encoding(path)
        with path.open("r", encoding=encoding, newline="") as stream:
            sample = stream.read(8192)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
                reader = csv.reader(stream, dialect)
            except csv.Error:
                reader = csv.reader(stream, delimiter=";")
            headers = [value.strip() for value in next(reader, [])]
            if headers_only:
                yield "CSV", headers, iter(())
            else:
                yield "CSV", headers, (
                    (number, tuple(row)) for number, row in enumerate(reader, start=2)
                )
        return
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                rows = worksheet.iter_rows(values_only=True)
                prefix = []
                for number, row in enumerate(rows, start=1):
                    prefix.append((number, tuple(row)))
                    if number >= 20:
                        break
                if not prefix:
                    continue
                header_position = _best_header_position(prefix)
                header_number, header_row = prefix[header_position]
                headers = _unique_headers(header_row)
                if headers_only:
                    yield worksheet.title, headers, iter(())
                    continue
                remaining_prefix = prefix[header_position + 1 :]
                remaining_rows = (
                    (number, tuple(row))
                    for number, row in enumerate(rows, start=len(prefix) + 1)
                )
                yield worksheet.title, headers, chain(remaining_prefix, remaining_rows)
        finally:
            workbook.close()
        return
    if suffix in {".xls", ".ods"}:
        engine = "xlrd" if suffix == ".xls" else "odf"
        excel = pd.ExcelFile(path, engine=engine)
        for sheet in excel.sheet_names:
            raw = pd.DataFrame(excel.parse(sheet_name=sheet, header=None))
            prefix = [(index + 1, tuple(row)) for index, row in enumerate(raw.head(20).itertuples(index=False, name=None))]
            if not prefix:
                continue
            header_position = _best_header_position(prefix)
            headers = _unique_headers(prefix[header_position][1])
            if headers_only:
                yield sheet, headers, iter(())
            else:
                rows = [
                    (index + header_position + 2, tuple(row))
                    for index, row in enumerate(raw.iloc[header_position + 1 :].itertuples(index=False, name=None))
                ]
                yield sheet, headers, iter(rows)
        return
    raise ValueError(f"Format non supporté : {suffix}")


def _best_header_position(rows: list[tuple[int, tuple[object, ...]]]) -> int:
    def score(row: tuple[object, ...]) -> int:
        return sum(3 if any(char.isalpha() for char in str(value)) else 1 for value in row if value is not None and str(value).strip())
    return max(range(len(rows)), key=lambda index: score(rows[index][1]))


def _unique_headers(row: tuple[object, ...]) -> list[str]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(row, start=1):
        base = str(value).strip() if value is not None and str(value).strip() else f"Colonne {index}"
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return headers


def _detect_encoding(path: Path) -> str:
    with path.open("rb") as stream:
        raw = stream.read(65536)
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            raw.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            pass
    return "latin-1"
