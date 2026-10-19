from __future__ import annotations

import csv
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Callable, Iterable, Iterator

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .models import DetectedTable, LoadCurvePoint
from .normalization import normalized_text


ProgressCallback = Callable[[str], None]

METER_ALIASES = (
    "numero du compteur", "numero compteur", "compteur", "numero prm", "prm",
    "id prm", "point de livraison", "pdl",
)
LABEL_ALIASES = ("libelle du prm", "libelle", "site", "nom du site", "information du prm")
TIMESTAMP_ALIASES = ("horodate", "datetime", "date heure", "dateheure", "timestamp")
DATE_ALIASES = ("date", "jour date")
TIME_ALIASES = ("heure", "time")
VALUE_ALIASES = ("valeur", "value", "puissance", "consommation")
UNIT_ALIASES = ("unite", "unit")
VARIABLE_ALIASES = ("nom de la variable", "variable")
YEAR_ALIASES = ("annee",)
MONTH_ALIASES = ("mois",)
DAY_ALIASES = ("jour",)
HOUR_ALIASES = ("heure",)
MINUTE_ALIASES = ("minute",)
SECOND_ALIASES = ("seconde", "secondes")


class LoadCurveError(ValueError):
    pass


def _alias_index(headers: list[str], aliases: Iterable[str]) -> int | None:
    normalized = [normalized_text(header).lstrip("# ") for header in headers]
    # L'ordre des alias exprime la priorité. Ainsi « Numéro du compteur »
    # est préféré à « #ID PRM » lorsque les deux colonnes existent.
    for alias in aliases:
        wanted = normalized_text(alias)
        for index, header in enumerate(normalized):
            if header == wanted:
                return index
    return None


def _column_map(headers: list[str]) -> dict[str, int | None]:
    return {
        "meter": _alias_index(headers, METER_ALIASES),
        "label": _alias_index(headers, LABEL_ALIASES),
        "timestamp": _alias_index(headers, TIMESTAMP_ALIASES),
        "date": _alias_index(headers, DATE_ALIASES),
        "time": _alias_index(headers, TIME_ALIASES),
        "value": _alias_index(headers, VALUE_ALIASES),
        "unit": _alias_index(headers, UNIT_ALIASES),
        "variable": _alias_index(headers, VARIABLE_ALIASES),
        "year": _alias_index(headers, YEAR_ALIASES),
        "month": _alias_index(headers, MONTH_ALIASES),
        "day": _alias_index(headers, DAY_ALIASES),
        "hour": _alias_index(headers, HOUR_ALIASES),
        "minute": _alias_index(headers, MINUTE_ALIASES),
        "second": _alias_index(headers, SECOND_ALIASES),
    }


def _looks_like_curve(headers: list[str]) -> bool:
    columns = _column_map(headers)
    has_components = all(columns[name] is not None for name in ("year", "month", "day", "hour"))
    has_datetime = columns["timestamp"] is not None or (
        columns["date"] is not None and columns["time"] is not None
    ) or has_components
    return columns["value"] is not None and has_datetime


def detect_xlsx_tables(path: str | Path) -> list[DetectedTable]:
    source = Path(path)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        detected = _detect_tables_in_workbook(source, workbook)
    finally:
        workbook.close()
    return detected


def _detect_tables_in_workbook(
    source: Path,
    workbook,
    progress: ProgressCallback | None = None,
) -> list[DetectedTable]:
    detected: list[DetectedTable] = []
    sheet_count = len(workbook.worksheets)
    for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
        if progress:
            progress(f"Détection : feuille {sheet_index}/{sheet_count} — {worksheet.title}")
        for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if row_number > 20:
                break
            headers = ["" if value is None else str(value).strip() for value in row]
            if _looks_like_curve(headers):
                detected.append(
                    DetectedTable(
                        source=source,
                        sheet=worksheet.title,
                        header_row=row_number,
                        columns=tuple(headers),
                        estimated_rows=max(0, worksheet.max_row - row_number),
                    )
                )
                break
    return detected


def load_curve_points(
    path: str | Path,
    *,
    progress: ProgressCallback | None = None,
) -> tuple[list[LoadCurvePoint], list[DetectedTable], list[str]]:
    source = Path(path)
    suffix = source.suffix.casefold()
    if suffix in {".xlsx", ".xlsm"}:
        return _load_xlsx(source, progress=progress)
    if suffix == ".csv":
        return _load_csv(source, progress=progress)
    raise LoadCurveError(f"Format non supporté : {source.suffix}")


def _load_xlsx(
    source: Path,
    *,
    progress: ProgressCallback | None,
) -> tuple[list[LoadCurvePoint], list[DetectedTable], list[str]]:
    if progress:
        progress(f"Étape 1/3 — Ouverture du classeur ({source.stat().st_size / 1_048_576:.1f} Mo)…")
    workbook = load_workbook(source, read_only=True, data_only=True)
    points: list[LoadCurvePoint] = []
    warnings: list[str] = []
    try:
        if progress:
            progress("Étape 1/3 — Recherche des feuilles de courbes de charge…")
        tables = _detect_tables_in_workbook(source, workbook, progress)
        if not tables:
            raise LoadCurveError("Aucune feuille contenant une courbe de charge n'a été détectée.")
        if progress:
            progress(f"Étape 2/3 — {len(tables)} feuille(s) trouvée(s). Lecture des données…")
        by_sheet = {table.sheet: table for table in tables}
        curve_index = 0
        for worksheet in workbook.worksheets:
            table = by_sheet.get(worksheet.title)
            if table is None:
                continue
            curve_index += 1
            if progress:
                progress(
                    f"Feuille {curve_index}/{len(tables)} — {worksheet.title} "
                    f"({table.estimated_rows:,} lignes estimées)".replace(",", " ")
                )
            points.extend(_points_from_worksheet(source, worksheet, table, warnings, progress))
    finally:
        workbook.close()
    if not points:
        raise LoadCurveError("Les feuilles détectées ne contiennent aucune mesure exploitable.")
    return points, tables, warnings


def _points_from_worksheet(
    source: Path,
    worksheet: Worksheet,
    table: DetectedTable,
    warnings: list[str],
    progress: ProgressCallback | None,
) -> Iterator[LoadCurvePoint]:
    columns = _column_map(list(table.columns))
    bad_rows = 0
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=table.header_row + 1, values_only=True),
        start=table.header_row + 1,
    ):
        completed_rows = row_number - table.header_row
        if progress and completed_rows > 0 and completed_rows % 2_500 == 0:
            percent = min(100, round(completed_rows * 100 / max(1, table.estimated_rows)))
            progress(
                f"  {worksheet.title} : {completed_rows:,}/{table.estimated_rows:,} lignes ({percent} %)"
                .replace(",", " ")
            )
        try:
            point = _point_from_row(row, columns, source.name, worksheet.title, row_number)
        except (TypeError, ValueError, OverflowError):
            bad_rows += 1
            continue
        if point is not None:
            yield point
    if bad_rows:
        warnings.append(f"{worksheet.title} : {bad_rows} ligne(s) invalide(s) ignorée(s).")
    if progress:
        progress(f"{worksheet.title} : lecture terminée.")


def _cell(row: tuple[object, ...] | list[object], index: int | None) -> object | None:
    return None if index is None or index >= len(row) else row[index]


def _point_from_row(
    row: tuple[object, ...] | list[object],
    columns: dict[str, int | None],
    source_name: str,
    sheet: str,
    row_number: int,
) -> LoadCurvePoint | None:
    raw_value = _cell(row, columns["value"])
    if raw_value is None or str(raw_value).strip() == "":
        return None
    variable = str(_cell(row, columns["variable"]) or "")
    if variable and "courbe de charge" not in normalized_text(variable):
        return None

    timestamp = _row_timestamp(row, columns)
    value = _number(raw_value)
    meter = _identifier(_cell(row, columns["meter"])) or sheet
    label = str(_cell(row, columns["label"]) or sheet).strip()
    unit = str(_cell(row, columns["unit"]) or "").strip()
    return LoadCurvePoint(meter, label, timestamp, value, unit, source_name, sheet, row_number)


def _row_timestamp(row: tuple[object, ...] | list[object], columns: dict[str, int | None]) -> datetime:
    component_names = ("year", "month", "day", "hour")
    if all(columns[name] is not None for name in component_names):
        return datetime(
            int(_number(_cell(row, columns["year"]))),
            int(_number(_cell(row, columns["month"]))),
            int(_number(_cell(row, columns["day"]))),
            int(_number(_cell(row, columns["hour"]))),
            int(_number(_cell(row, columns["minute"]) or 0)),
            int(_number(_cell(row, columns["second"]) or 0)),
        )
    raw_timestamp = _cell(row, columns["timestamp"])
    if raw_timestamp is not None:
        return _datetime_value(raw_timestamp)
    raw_date = _cell(row, columns["date"])
    raw_time = _cell(row, columns["time"])
    return _combine_date_time(raw_date, raw_time)


def _datetime_value(value: object) -> datetime:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, time())
    if isinstance(value, (int, float)):
        return datetime(1899, 12, 30) + timedelta(days=float(value))
    text = str(value).strip()
    for fmt in (
        "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%SZ", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    raise ValueError(f"Horodate invalide : {text}")


def _combine_date_time(date_value: object, time_value: object) -> datetime:
    if isinstance(date_value, datetime):
        base_date = date_value.date()
    elif isinstance(date_value, date):
        base_date = date_value
    else:
        text = str(date_value).strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                base_date = datetime.strptime(text, fmt).date()
                break
            except ValueError:
                continue
        else:
            raise ValueError(f"Date invalide : {text}")
    if isinstance(time_value, datetime):
        base_time = time_value.time()
    elif isinstance(time_value, time):
        base_time = time_value
    else:
        text = str(time_value).strip()
        for fmt in ("%H:%M:%S", "%H:%M"):
            try:
                base_time = datetime.strptime(text, fmt).time()
                break
            except ValueError:
                continue
        else:
            raise ValueError(f"Heure invalide : {text}")
    return datetime.combine(base_date, base_time)


def _number(value: object) -> float:
    if isinstance(value, bool):
        raise ValueError("Valeur booléenne")
    if isinstance(value, (int, float)):
        return float(value)
    return float(str(value).strip().replace(" ", "").replace(",", "."))


def _identifier(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _load_csv(
    source: Path,
    *,
    progress: ProgressCallback | None,
) -> tuple[list[LoadCurvePoint], list[DetectedTable], list[str]]:
    if progress:
        progress(f"Lecture de {source.name}…")
    encoding = _detect_encoding(source)
    with source.open("r", encoding=encoding, newline="") as stream:
        sample = stream.read(8192)
        stream.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";"
        reader = csv.reader(stream, dialect)
        try:
            headers = [value.strip() for value in next(reader)]
        except StopIteration as exc:
            raise LoadCurveError("Le CSV est vide.") from exc
        if not _looks_like_curve(headers):
            raise LoadCurveError("Le CSV ne contient pas les colonnes date/heure/valeur attendues.")
        columns = _column_map(headers)
        points: list[LoadCurvePoint] = []
        bad_rows = 0
        for row_number, row in enumerate(reader, start=2):
            try:
                point = _point_from_row(row, columns, source.name, "CSV", row_number)
            except (TypeError, ValueError, OverflowError):
                bad_rows += 1
                continue
            if point is not None:
                points.append(point)
    table = DetectedTable(source, "CSV", 1, tuple(headers), len(points) + bad_rows)
    warnings = [f"CSV : {bad_rows} ligne(s) invalide(s) ignorée(s)."] if bad_rows else []
    return points, [table], warnings


def _detect_encoding(path: Path) -> str:
    with path.open("rb") as stream:
        raw = stream.read(65536)
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            raw.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    return "latin-1"

