from __future__ import annotations

from datetime import datetime

from openpyxl import Workbook

from autocalsol_tools.loader import detect_xlsx_tables, load_curve_points


def test_detects_only_load_curve_sheet(tmp_path):
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    cover = workbook.active
    cover.title = "Accueil"
    cover.append(["Titre", "Rapport"])
    curve = workbook.create_sheet("Compteur A")
    curve.append(["Numéro du compteur", "Horodate", "Unité", "Valeur", "Nom de la variable"])
    curve.append(["00123", datetime(2025, 1, 1), "Watt", 1200, "B01_Courbe de charge Consommation"])
    workbook.save(path)

    tables = detect_xlsx_tables(path)
    points, _, warnings = load_curve_points(path)

    assert [table.sheet for table in tables] == ["Compteur A"]
    assert len(points) == 1
    assert points[0].meter_id == "00123"
    assert warnings == []


def test_prefers_meter_number_over_internal_prm_id(tmp_path):
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["#ID PRM", "Numéro du compteur", "Horodate", "Valeur"])
    sheet.append([881453, "04207607671840", datetime(2025, 1, 1), 1200])
    workbook.save(path)

    points, _, _ = load_curve_points(path)
    assert points[0].meter_id == "04207607671840"


def test_loads_legacy_csv(tmp_path):
    path = tmp_path / "curve.csv"
    path.write_text("date;heure;valeur\n01/01/2025;00:00:00;1,5\n", encoding="utf-8")
    points, tables, warnings = load_curve_points(path)
    assert len(points) == 1
    assert points[0].timestamp == datetime(2025, 1, 1)
    assert points[0].value == 1.5
    assert tables[0].sheet == "CSV"
    assert warnings == []

